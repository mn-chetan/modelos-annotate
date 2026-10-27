from __future__ import annotations

import io
import mimetypes
import re
from pathlib import Path
from typing import Iterable

import pymupdf as fitz
from docx import Document
from fastapi import UploadFile
from pydantic import BaseModel

from app.core.config import get_settings
from app.core.exceptions import AppError
from app.schemas.source import PageText, SourceDocument
from app.services.ocr_service import OCRService


class TextExtractionService:
    def __init__(self, ocr_service: OCRService):
        self.settings = get_settings()
        self.ocr_service = ocr_service

    @staticmethod
    def _normalize_whitespace(text: str) -> str:
        return re.sub(r"\s+", " ", text).strip()

    async def extract_from_text(self, text: str) -> SourceDocument:
        normalized = self._normalize_whitespace(text)
        if not normalized:
            raise AppError("EMPTY_TEXT", "No text was provided.", status_code=422)
        return SourceDocument(source_type="TEXT", text=normalized, extracted_text=normalized)

    async def extract_from_upload(self, upload_file: UploadFile) -> SourceDocument:
        content = await upload_file.read()
        try:
            if len(content) > self.settings.max_upload_size_bytes:
                raise AppError("FILE_TOO_LARGE", "Uploaded file exceeds the allowed size.", status_code=413)
            filename = upload_file.filename or "uploaded_file"
            mime_type = upload_file.content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
            source_type = self._infer_source_type(filename, mime_type)
            if source_type == "PDF":
                return await self._extract_pdf(content, filename, mime_type)
            if source_type == "DOCX":
                return await self._extract_docx(content, filename, mime_type)
            if source_type == "TXT":
                return await self._extract_txt(content, filename, mime_type)
            if source_type == "IMAGE":
                return await self._extract_image(content, filename, mime_type)
            raise AppError("UNSUPPORTED_FILE_TYPE", f"Unsupported file type: {mime_type}", status_code=415)
        finally:
            await upload_file.close()

    @staticmethod
    def _infer_source_type(filename: str, mime_type: str) -> str:
        extension = Path(filename).suffix.lower()
        if mime_type == "application/pdf" or extension == ".pdf":
            return "PDF"
        if mime_type in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"} or extension == ".docx":
            return "DOCX"
        if mime_type.startswith("text/") or extension == ".txt":
            return "TXT"
        if mime_type.startswith("image/") or extension in {".jpg", ".jpeg", ".png"}:
            return "IMAGE"
        return "UNKNOWN"

    async def _extract_pdf(self, content: bytes, filename: str, mime_type: str) -> SourceDocument:
        try:
            pdf = fitz.open(stream=content, filetype="pdf")
        except Exception as exc:
            raise AppError("INVALID_PDF", "The uploaded PDF could not be read.", status_code=422) from exc

        pages: list[PageText] = []
        text_segments: list[str] = []
        for index, page in enumerate(pdf, start=1):
            page_text = self._normalize_whitespace(page.get_text("text"))
            if page_text:
                pages.append(PageText(page=index, text=page_text))
                text_segments.append(page_text)
            else:
                rendered = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                ocr_text = self._normalize_whitespace(await self.ocr_service.extract_text(rendered.tobytes("png")))
                pages.append(PageText(page=index, text=ocr_text))
                text_segments.append(ocr_text)
        extracted_text = self._normalize_whitespace("\n".join(text_segments))
        if not extracted_text:
            raise AppError("EMPTY_EXTRACTED_TEXT", "No text could be extracted from the PDF.", status_code=422)
        return SourceDocument(
            source_type="PDF",
            original_filename=filename,
            mime_type=mime_type,
            file_size=len(content),
            extracted_text=extracted_text,
            pages=pages,
        )

    async def _extract_docx(self, content: bytes, filename: str, mime_type: str) -> SourceDocument:
        try:
            document = Document(io.BytesIO(content))
        except Exception as exc:
            raise AppError("INVALID_DOCX", "The uploaded DOCX could not be read.", status_code=422) from exc

        paragraphs = [self._normalize_whitespace(paragraph.text) for paragraph in document.paragraphs if paragraph.text.strip()]
        table_cells: list[str] = []
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(self._normalize_whitespace(cell.text) for cell in row.cells if cell.text.strip())
                if row_text:
                    table_cells.append(row_text)
        segments = [segment for segment in paragraphs + table_cells if segment]
        extracted_text = self._normalize_whitespace("\n".join(segments))
        if not extracted_text:
            raise AppError("EMPTY_EXTRACTED_TEXT", "No text could be extracted from the DOCX.", status_code=422)
        return SourceDocument(
            source_type="DOCX",
            original_filename=filename,
            mime_type=mime_type,
            file_size=len(content),
            extracted_text=extracted_text,
            pages=[PageText(page=1, text=extracted_text)],
        )

    async def _extract_txt(self, content: bytes, filename: str, mime_type: str) -> SourceDocument:
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise AppError("INVALID_TXT", "The uploaded TXT could not be decoded.", status_code=422)

        normalized = self._normalize_whitespace(text)
        if not normalized:
            raise AppError("EMPTY_EXTRACTED_TEXT", "No text could be extracted from the TXT file.", status_code=422)
        return SourceDocument(
            source_type="TXT",
            original_filename=filename,
            mime_type=mime_type,
            file_size=len(content),
            extracted_text=normalized,
            pages=[PageText(page=1, text=normalized)],
        )

    async def _extract_image(self, content: bytes, filename: str, mime_type: str) -> SourceDocument:
        text = self._normalize_whitespace(await self.ocr_service.extract_text(content))
        if not text:
            raise AppError("EMPTY_EXTRACTED_TEXT", "No text could be extracted from the image.", status_code=422)
        return SourceDocument(
            source_type="IMAGE",
            original_filename=filename,
            mime_type=mime_type,
            file_size=len(content),
            extracted_text=text,
            pages=[PageText(page=1, text=text)],
        )
