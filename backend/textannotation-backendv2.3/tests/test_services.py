from __future__ import annotations

import io

import fitz
import pytest
from docx import Document
from PIL import Image, ImageDraw

from app.core.exceptions import AppError
from app.core.config import Settings
from app.schemas.analysis import GeminiRelationExtractionResponse, GeminiTextClassificationResponse, TextClassificationAnalysisResult, RelationExtractionAnalysisResult, RelationExtractionRelation
from app.schemas.entity import EntityAnalysisResult, EntityMatch, EntityValidationResponse
from app.schemas.sentiment import GeminiSentimentResponse, SentimentAnalysisResult
from app.schemas.source import PageText, SourceDocument
from app.services.drug_matching_service import DrugMatchingService
from app.services.gemini_service import GeminiService
from app.services.relation_extraction_service import RelationExtractionService
from app.services.ocr_service import OCRService
from app.services.sentiment_service import SentimentService
from app.services.text_classification_service import TextClassificationService
from app.services.text_extraction_service import TextExtractionService


def test_entity_extraction_prompt_mentions_case_insensitive_semantic_matching():
    from app.prompts.drug_validation_prompt import build_entity_extraction_prompt

    prompt = build_entity_extraction_prompt(["Python", "Software Engineer"], "Python developer")

    assert "single document" in prompt
    assert "case-insensitively" in prompt
    assert "close semantic variations" in prompt
    assert "Return strict JSON with key: entities" in prompt


def test_text_classification_prompt_mentions_issue_labels():
    from app.prompts.text_classification_prompt import build_text_classification_prompt

    prompt = build_text_classification_prompt("Unable to login after password reset")

    assert "text classification only" in prompt
    assert "Login Issue" in prompt
    assert "Return strict JSON with keys: analysis_type, label, confidence, reason" in prompt


def test_relation_extraction_prompt_mentions_contract_relations():
    from app.prompts.relation_extraction_prompt import build_relation_extraction_prompt

    prompt = build_relation_extraction_prompt("John Smith joined Tech Mahindra as a Senior Developer.")

    assert "relation extraction only" in prompt
    assert "contract-style relationships" in prompt
    assert "RELATION_EXTRACTION" in prompt
    assert "WORKS_FOR" in prompt


class FakeOCRService(OCRService):
    def __init__(self, text: str = "Scanned PDF text"):
        self.text = text

    async def extract_text(self, image_bytes: bytes) -> str:
        return self.text


class FakeGeminiService:
    def __init__(self, response):
        self.response = response
        self.prompts: list[str] = []

    async def generate_structured(self, prompt: str, response_model):
        self.prompts.append(prompt)
        if isinstance(self.response, Exception):
            raise self.response
        return self.response


class FakeUploadFile:
    def __init__(self, filename: str, content: bytes, content_type: str):
        self.filename = filename
        self._content = content
        self.content_type = content_type
        self.closed = False

    async def read(self) -> bytes:
        return self._content

    async def close(self) -> None:
        self.closed = True


class FakeGoogleCredentials:
    def __init__(self):
        self.expired = False
        self.refresh_token = "token"
        self.token = "access-token"
        self.project_id = None

    def refresh(self, request):
        return None


class ExpiredGoogleCredentials(FakeGoogleCredentials):
    def refresh(self, request):
        from google.auth.exceptions import RefreshError

        raise RefreshError("Reauthentication is needed")


@pytest.mark.asyncio
async def test_direct_text_extraction_normalizes_whitespace():
    service = TextExtractionService(FakeOCRService())
    source = await service.extract_from_text("The medicine   worked well.\n\nPackaging was good.")
    assert source.source_type == "TEXT"
    assert source.extracted_text == "The medicine worked well. Packaging was good."


@pytest.mark.asyncio
async def test_txt_docx_pdf_and_image_extraction():
    service = TextExtractionService(FakeOCRService())

    txt_file = FakeUploadFile("notes.txt", b"Paracetamol 500 mg", "text/plain")
    txt_source = await service.extract_from_upload(txt_file)
    assert txt_source.source_type == "TXT"
    assert "Paracetamol" in txt_source.extracted_text

    docx_buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Paracetamol 500 mg")
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Caffeine 30 mg"
    document.save(docx_buffer)
    docx_buffer.seek(0)
    docx_file = FakeUploadFile(
        "ingredients.docx",
        docx_buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    docx_source = await service.extract_from_upload(docx_file)
    assert docx_source.source_type == "DOCX"
    assert "Caffeine 30 mg" in docx_source.extracted_text

    pdf_doc = fitz.open()
    pdf_page = pdf_doc.new_page()
    pdf_page.insert_text((72, 72), "Ibuprofen 200 mg")
    pdf_bytes = pdf_doc.tobytes()
    pdf_file = FakeUploadFile("sample.pdf", pdf_bytes, "application/pdf")
    pdf_source = await service.extract_from_upload(pdf_file)
    assert pdf_source.source_type == "PDF"
    assert pdf_source.pages[0].page == 1
    assert "Ibuprofen" in pdf_source.extracted_text

    image = Image.new("RGB", (240, 80), color="white")
    draw = ImageDraw.Draw(image)
    draw.text((10, 20), "Paracetamol", fill="black")
    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")
    image_buffer.seek(0)
    image_file = FakeUploadFile("label.png", image_buffer.getvalue(), "image/png")
    image_source = await service.extract_from_upload(image_file)
    assert image_source.source_type == "IMAGE"
    assert image_source.extracted_text == "Scanned PDF text"


@pytest.mark.asyncio
async def test_exact_drug_match():
    service = DrugMatchingService(
        FakeGeminiService(
            EntityValidationResponse(
                entities=[EntityMatch(entity_name="Paracetamol", matched_text="Paracetamol", tag="OTHER", found=True, confidence_score=0.99)]
            )
        )
    )
    source = SourceDocument(source_type="TEXT", extracted_text="Each tablet contains Paracetamol 500 mg.", pages=[PageText(page=1, text="Each tablet contains Paracetamol 500 mg.")])
    result = await service.analyze(source, "Paracetamol")
    assert result.entities[0].entity_name == "Paracetamol"
    assert result.entities[0].found is True


@pytest.mark.asyncio
async def test_drug_not_found():
    service = DrugMatchingService(
        FakeGeminiService(
            EntityValidationResponse(
                entities=[EntityMatch(entity_name="Ibuprofen", matched_text=None, tag="OTHER", found=False, confidence_score=0.12)]
            )
        )
    )
    source = SourceDocument(source_type="TEXT", extracted_text="Each tablet contains Paracetamol 500 mg.", pages=[PageText(page=1, text="Each tablet contains Paracetamol 500 mg.")])
    result = await service.analyze(source, "Ibuprofen")
    assert result.entities[0].entity_name == "Ibuprofen"
    assert result.entities[0].found is False


@pytest.mark.asyncio
async def test_fuzzy_and_semantic_drug_match(monkeypatch):
    semantic_response = EntityValidationResponse(entities=[EntityMatch(entity_name="Paracetamol", matched_text="Paracetamol 500 mg", tag="OTHER", found=True, confidence_score=0.87)])
    fake_gemini = FakeGeminiService(semantic_response)
    service = DrugMatchingService(fake_gemini)
    source = SourceDocument(source_type="TEXT", extracted_text="Paracetamol 500 mg", pages=[PageText(page=1, text="Paracetamol 500 mg")])

    result = await service.analyze(source, "Paracetamol, Ibuprofen")
    assert result.entities[0].found is True
    assert fake_gemini.prompts


@pytest.mark.asyncio
async def test_positive_negative_neutral_mixed_sentiment():
    positive = GeminiSentimentResponse(sentiment="POSITIVE", confidence=0.95, reason="Positive")
    negative = GeminiSentimentResponse(sentiment="NEGATIVE", confidence=0.95, reason="Negative")
    neutral = GeminiSentimentResponse(sentiment="NEUTRAL", confidence=0.90, reason="Neutral")
    mixed = GeminiSentimentResponse(sentiment="MIXED", confidence=0.90, reason="Mixed")

    service = SentimentService(FakeGeminiService(positive))
    source = SourceDocument(source_type="TEXT", extracted_text="The medicine worked well.", pages=[PageText(page=1, text="The medicine worked well.")])
    result = await service.analyze(source)
    assert result.sentiment == "POSITIVE"

    service = SentimentService(FakeGeminiService(negative))
    source = SourceDocument(source_type="TEXT", extracted_text="The packaging was damaged.", pages=[PageText(page=1, text="The packaging was damaged.")])
    result = await service.analyze(source)
    assert result.sentiment == "NEGATIVE"

    service = SentimentService(FakeGeminiService(neutral))
    source = SourceDocument(source_type="TEXT", extracted_text="The box arrived today.", pages=[PageText(page=1, text="The box arrived today.")])
    result = await service.analyze(source)
    assert result.sentiment == "NEUTRAL"

    service = SentimentService(FakeGeminiService(mixed))
    source = SourceDocument(source_type="TEXT", extracted_text="The medicine worked well but the packaging was damaged.", pages=[PageText(page=1, text="The medicine worked well but the packaging was damaged.")])
    result = await service.analyze(source)
    assert result.sentiment == "MIXED"


@pytest.mark.asyncio
async def test_text_classification_uses_chunked_gemini_results():
    response = GeminiTextClassificationResponse(
        results=[
            GeminiTextClassificationResponse.ClassificationItem(
                text="Unable to login after password reset.",
                category="Login Issue",
                confidence=0.96,
                reasoning="Mentions login after reset.",
            )
        ]
    )
    service = TextClassificationService(FakeGeminiService(response))
    source = SourceDocument(
        source_type="TEXT",
        extracted_text="Unable to login after password reset.\n\nApplication takes 30 seconds to load.",
        pages=[PageText(page=1, text="Unable to login after password reset.")],
    )

    result = await service.analyze(source)

    assert isinstance(result, TextClassificationAnalysisResult)
    assert result.label == "Login Issue"


@pytest.mark.asyncio
async def test_relation_extraction_uses_chunked_gemini_results():
    response = GeminiRelationExtractionResponse(
        relations=[
            GeminiRelationExtractionResponse.RelationItem(
                source="John Smith",
                relation="WORKS_FOR",
                target="Tech Mahindra",
                confidence=0.97,
                reasoning="Employment relation is explicit.",
            )
        ]
    )
    service = RelationExtractionService(FakeGeminiService(response))
    source = SourceDocument(
        source_type="TEXT",
        extracted_text="John Smith joined Tech Mahindra as a Senior Developer.",
        pages=[PageText(page=1, text="John Smith joined Tech Mahindra as a Senior Developer.")],
    )

    result = await service.analyze(source)

    assert isinstance(result, RelationExtractionAnalysisResult)
    assert result.relations[0].relation == "WORKS_FOR"
    assert result.relations[0].source == "John Smith"
    assert result.reason == "Extracted 1 relation(s) from the provided text."


@pytest.mark.asyncio
async def test_text_classification_splits_numbered_list_items():
    response = GeminiTextClassificationResponse(
        results=[
            GeminiTextClassificationResponse.ClassificationItem(
                text="Unable to login after password reset.",
                category="Login Issue",
                confidence=0.94,
                reasoning="Login failure after reset.",
            )
        ]
    )
    service = TextClassificationService(FakeGeminiService(response))
    source = SourceDocument(
        source_type="PDF",
        extracted_text="1. Unable to login after password reset 2. Application takes 30 seconds to load 3. Payment deducted but order not created",
        pages=[PageText(page=1, text="1. Unable to login after password reset 2. Application takes 30 seconds to load 3. Payment deducted but order not created")],
    )

    result = await service.analyze(source)

    assert len(result.chunk_results) == 3
    assert result.chunk_results[0].chunk_text == "Unable to login after password reset"
    assert result.chunk_results[1].chunk_text == "Application takes 30 seconds to load"
    assert result.chunk_results[2].chunk_text == "Payment deducted but order not created"


@pytest.mark.asyncio
async def test_text_classification_splits_bullet_list_items():
    response = GeminiTextClassificationResponse(
        results=[
            GeminiTextClassificationResponse.ClassificationItem(
                text="Unable to login after password reset",
                category="Login Issue",
                confidence=0.94,
                reasoning="Login failure after reset.",
            )
        ]
    )
    service = TextClassificationService(FakeGeminiService(response))
    source = SourceDocument(
        source_type="PDF",
        extracted_text="• Unable to login after password reset • Application takes 30 seconds to load • Payment deducted but order not created",
        pages=[PageText(page=1, text="• Unable to login after password reset • Application takes 30 seconds to load • Payment deducted but order not created")],
    )

    result = await service.analyze(source)

    assert len(result.chunk_results) == 3
    assert result.chunk_results[0].chunk_text == "Unable to login after password reset"
    assert result.chunk_results[1].chunk_text == "Application takes 30 seconds to load"
    assert result.chunk_results[2].chunk_text == "Payment deducted but order not created"


@pytest.mark.asyncio
async def test_gemini_failure_raises_app_error():
    service = SentimentService(FakeGeminiService(AppError("GEMINI_FAILURE", "failed", 502)))
    source = SourceDocument(source_type="TEXT", extracted_text="The medicine worked well.", pages=[PageText(page=1, text="The medicine worked well.")])
    with pytest.raises(AppError):
        await service.analyze(source)


@pytest.mark.asyncio
async def test_gemini_service_uses_project_id_from_adc(monkeypatch):
    fake_credentials = FakeGoogleCredentials()

    async def fake_to_thread(func, *args, **kwargs):
        return fake_credentials, "adc-project"

    monkeypatch.setattr("app.services.gemini_service.asyncio.to_thread", fake_to_thread)

    service = GeminiService(project_id=None)
    await service._load_credentials()

    assert service.project_id == "adc-project"
    assert service._credentials is fake_credentials


@pytest.mark.asyncio
async def test_gemini_service_refresh_failure_raises_auth_error(monkeypatch):
    expired_credentials = ExpiredGoogleCredentials()

    async def fake_to_thread(func, *args, **kwargs):
        return expired_credentials, "adc-project"

    monkeypatch.setattr("app.services.gemini_service.asyncio.to_thread", fake_to_thread)

    service = GeminiService(project_id="adc-project")

    with pytest.raises(AppError) as exc_info:
        await service._refresh_headers()

    assert exc_info.value.code == "GOOGLE_CLOUD_AUTH_REAUTH_REQUIRED"


def test_gemini_default_model_targets_current_vertex_model():
    settings = Settings()
    service = GeminiService(project_id="test-project", location="us-central1")

    assert settings.gemini_model == "gemini-2.5-flash"
    assert service._endpoint().endswith("/publishers/google/models/gemini-2.5-flash:generateContent")
